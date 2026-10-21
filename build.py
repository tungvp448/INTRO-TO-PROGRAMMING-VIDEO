from pathlib import Path
import io, contextlib, json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
slides = []
def slide(title, seconds, bullets, code, output, explanation, visual, error=None):
    slides.append(dict(title=title, seconds=seconds, bullets=bullets, code=code, output=output, explanation=explanation, visual=visual, error=error))

slide('Learning objectives', 40,
['Explain variables, assignment and basic Python data types.', 'Use sequence indexing and slicing.', 'Create, update and copy lists.', 'Combine these ideas in a small student score program.'],
'student_name = "Alex"\nscores = [8.5, 7.0, 9.0]\nprint(student_name)\nprint(scores)',
'Alex\n[8.5, 7.0, 9.0]',
'''This presentation covers two topics: Variables and Data Types, and Lists and Sequences. We begin with how Python associates names with data, then explain how data types affect the operations we can perform. Next, we introduce sequences and examine lists in more detail. Finally, we combine these concepts in a small program that stores and updates student scores. The example on this slide shows the destination of the lesson: one name refers to a student's name, while another refers to a collection of scores. We will explain the syntax before developing the complete program.''',
'Show the two topic names and the four learning objectives. Keep the opening code as a small preview, without explaining every operation yet.')

slide('Variables and assignment', 70,
['A variable name refers to an object.', 'Assignment evaluates the right side and binds the name on the left.', 'Reassignment changes what a name refers to.', '= assigns; == compares values.'],
'score = 8\nscore = score + 1\nprint(score)\nprint(score == 9)', '9\nTrue',
'''A variable is a name that refers to an object. In score equals eight, Python evaluates the value on the right and binds the name score to that integer object. The assignment operator is not a mathematical statement that both sides must remain equal forever. In score equals score plus one, Python first reads the current value, calculates nine, and then rebinds score to the result. The original integer object is not changed into another integer. The final line uses two equals signs to compare the value of score with nine. This comparison produces True. Using a name before it has been defined normally raises a NameError. For our examples, we define every name before using it.''',
'Use two simple states: score refers to 8, then score refers to 9. Label the second state reassignment rather than changing the integer object.')

slide('Naming variables clearly', 50,
['Use descriptive names such as student_name and average_score.', 'For simple ASCII names: begin with a letter or underscore; continue with letters, digits or underscores.', 'Names are case sensitive; keywords cannot be variable names.', 'Use snake_case by convention; avoid replacing built-in names such as list or str.'],
'student_name = "Alex"\nscore_1 = 8.5\nScore_1 = 9.0\nprint(score_1, Score_1)', '8.5 9.0',
'''Good names describe the role of a value. Student underscore name is clearer than an unrelated single letter when storing a student's name. For the simple ASCII names used here, a name begins with a letter or underscore and may then contain digits. Spaces and hyphens cannot form part of such a name. Python also permits many Unicode letters, but simple English names are sufficient for this lesson. Names are case sensitive, so the two score names shown here are different. Reserved keywords, such as class and for, cannot be used as variable names. Snake case is a readability convention rather than a syntax requirement. Built-in names such as list are technically assignable, but replacing them makes their original functions harder to use.''',
'Place valid examples beside invalid examples: 1score, student-name, and class. Label list = 3 as legal but discouraged, not a syntax error.')

slide('Basic data types', 80,
['int: integers, such as 18.', 'float: floating-point numbers, such as 8.5.', 'str: text, such as "Alex".', 'bool: True or False.', 'None: absence of a value; its type is NoneType.'],
'age = 18\nscore = 8.5\nstudent_name = "Alex"\npassed = True\nfeedback = None\nprint(type(age))\nprint(type(score))\nprint(type(student_name))\nprint(type(passed))\nprint(type(feedback))',
"<class 'int'>\n<class 'float'>\n<class 'str'>\n<class 'bool'>\n<class 'NoneType'>",
'''A type describes an object's category and helps determine which operations it supports. Integers represent whole numbers, including negative numbers and zero. Floats represent floating-point numbers; many decimal fractions are approximations in this representation. Strings represent text and are written with matching quotation marks. A boolean has one of two values, True or False, with capital initial letters. None is a special value commonly used to represent missing or not-yet-available information. It is not the string None, an empty string, or the number zero. The built-in type function reports the type of an object. These are the basic types needed for our lesson, not an exhaustive list of Python types. Lists and other sequence types will be introduced shortly.''',
'Use a five-row type table on the actual slide. Show the type outputs beside their values rather than displaying two tall code columns.')

slide('Dynamic typing and operations', 70,
['Objects have types; a name can be rebound to an object of another type.', 'The meaning of an operation depends on the operand types.', 'Python does not automatically convert a numeric string for addition.'],
'value = 10\nprint(type(value))\nvalue = "ten"\nprint(type(value))\nprint(2 + 3)\nprint("2" + "3")\nprint(8 / 2)\nprint(8.5 >= 5.0)',
"<class 'int'>\n<class 'str'>\n5\n23\n4.0\nTrue",
'''Python is dynamically typed: we do not declare a fixed type for these variable names before assignment. The name value first refers to an integer and later to a string. This does not mean that Python has no types, or that the original integer has become text. Types still control operations. Adding two integers performs arithmetic, while adding two strings concatenates their text. The printed result twenty-three in this example is a string, even though print does not display its quotation marks. For the built-in integers shown here, ordinary division produces a float. A numerical comparison produces a boolean. Combining a string and an integer with plus, such as the string two plus the integer three, raises a TypeError rather than silently choosing a conversion.''',
'Pair the numeric addition and string concatenation examples. Label the result types explicitly.')

slide('Converting between types', 75,
['float("8.5") converts suitable text to a float.', 'int("18") converts integer text; int(8.9) truncates toward zero.', 'str(18) produces text.', 'input() returns a string; convert it before numerical calculation.'],
'raw_score = "8.5"\nscore = float(raw_score)\nprint(score + 1)\nprint(int("18"))\nprint(int(8.9))\nprint(str(18) + " years")',
'9.5\n18\n8\n18 years',
'''Conversion creates a suitable value for the operation we intend to perform. Here, raw score contains text. Float converts that text to the number eight point five, so adding one produces nine point five. Int can parse text representing an integer, but int applied directly to the string eight point five raises a ValueError. That text is not an integer representation accepted by int. Applied to a float such as eight point nine, int truncates toward zero rather than rounding to the nearest integer. Str converts a value to text, allowing us to concatenate it with other text. This matters when using input, because input returns a string even when a person types digits. Conversion can fail, so the example assumes that the score text is valid.''',
'Show a short conversion path: text "8.5", then float(...), then number 8.5. Keep keyboard input as a concept, not an interactive recording step.')

slide('Sequences and their main types', 75,
['A sequence contains items accessible by position.', 'Ordered means positional order, not sorted order.', 'list: mutable sequence; tuple: immutable sequence.', 'str: immutable text sequence; range: immutable integer sequence.'],
'scores = [8.5, 7.0, 9.0]\nfixed_scores = (8.5, 7.0, 9.0)\ncourse = "Python"\npositions = range(3)\nprint(list(positions))\nprint(len(course))', ' [0, 1, 2]\n6'.lstrip(),
'''When we need several values in a defined order, a sequence provides access by position. Ordered does not mean sorted: our scores remain in the order eight point five, seven, and nine. A list uses square brackets and allows changes to its contents. A tuple is another sequence, commonly written with parentheses, whose item references cannot be replaced after creation. A string is an immutable sequence of Unicode characters. Range represents an arithmetic sequence of integers without storing every element as a list. Range of three represents zero, one, and two; the stop value three is excluded. We convert it to a list here only to display those values. These types share useful sequence operations, but they do not all support every list method. We will focus on the shared operations first.''',
'Use a comparison table with type, example, and mutable columns. Put range in a brief fourth row; do not introduce loops here.')

slide('Indexing by position', 65,
['Indexing selects one item: sequence[index].', 'The first item has index 0.', 'Negative indices count from the end; -1 selects the last item.', 'A missing position raises IndexError.'],
'scores = [8.5, 7.0, 9.0, 6.5]\nprint(scores[0])\nprint(scores[2])\nprint(scores[-1])\nprint(scores[-4])', '8.5\n9.0\n6.5\n8.5',
'''Indexing retrieves one item from a sequence. Python uses zero-based indexing, so a four-item list has nonnegative indices zero through three. Scores at index zero gives the first score, while index two gives the third score. Negative indexing starts at the end: minus one selects the last item and minus four selects the first item in this particular list. For a sequence of length n, valid nonnegative indices run from zero to n minus one. A length of four does not make index four valid. Trying to access that index raises an IndexError. The same positional idea applies to strings and tuples. Indexing an empty list also raises an IndexError because there is no item to retrieve.''',
'Draw four aligned cells with values 8.5, 7.0, 9.0, 6.5; show indices 0 to 3 above and -4 to -1 below.')

slide('Slicing a sequence', 80,
['sequence[start:stop:step] selects a subsequence.', 'With a positive step, start is included and stop is excluded.', 'Omitted bounds with a positive step use the beginning and end.', 'For lists, slicing creates a new list; it is a shallow copy.'],
'scores = [8.5, 7.0, 9.0, 6.5]\nprint(scores[1:3])\nprint(scores[:2])\nprint(scores[-2:])\nprint(scores[::2])\nprint(scores[::-1])\nprint(scores[1:99])',
'[7.0, 9.0]\n[8.5, 7.0]\n[9.0, 6.5]\n[8.5, 9.0]\n[6.5, 9.0, 7.0, 8.5]\n[7.0, 9.0, 6.5]',
'''Slicing selects a subsequence rather than one item. With the usual positive step, start is included and stop is excluded. Scores from one to three therefore contains the items at indices one and two. Omitting start begins at the beginning, and omitting stop continues to the end. The default step is one. A step of two selects every second item. The special form with both bounds omitted and a step of minus one traverses the sequence backwards. A step of zero is invalid. Unlike indexing, slice bounds beyond the available positions are clipped, so slicing from one to ninety-nine does not raise an IndexError here. A list slice creates a new list without changing the original. Its elements are not recursively copied, a distinction we will revisit when discussing copying.''',
'Emphasize [1:3] with highlighted cells. Reveal the step and reverse examples afterward so the first explanation remains simple.')

slide('Common sequence operations', 60,
['len(sequence) returns its length.', 'in and not in test membership.', 'Compatible lists, tuples and strings support concatenation and repetition.', 'Indexing, slicing and membership also work with range; range does not support + or *.'],
'scores = [8.5, 7.0, 9.0]\nprint(len(scores))\nprint(7.0 in scores)\nprint(10.0 not in scores)\nprint([1, 2] + [3])\nprint([0] * 3)\nprint("Py" in "Python")',
'3\nTrue\nTrue\n[1, 2, 3]\n[0, 0, 0]\nTrue',
'''Several operations make sequences convenient to inspect and combine. Len gives the number of items, and membership checks produce boolean results. In a list, membership checks whether an equal item is present. In a string, in can test whether a substring occurs. Plus concatenates the two lists shown here; it does not add corresponding numbers. Multiplying this one-item list by three repeats its contents. These operations create new lists in these examples. Strings and tuples also support their own compatible concatenation and repetition, but range does not support plus or multiplication. We should therefore avoid assuming that every sequence supports every operation. Also, repetition repeats references: repeating a nested mutable list can introduce shared objects, so we use immutable integer elements in this introductory example.''',
'Arrange the examples by purpose: measure, test, combine. Keep the nested-list caution in the explanation rather than adding another example to the slide.')

slide('Creating and updating lists', 65,
['Lists are ordered, mutable sequences.', 'They can be empty, contain duplicates, and contain different types.', 'Assigning to an existing index replaces one item.', 'A consistent element type often makes calculations simpler.'],
'empty = []\nrecord = ["Alex", 18, True]\nscores = [8.5, 7.0, 9.0]\nscores[1] = 7.5\nprint(scores)\nprint(len(empty))', '[8.5, 7.5, 9.0]\n0',
'''A list groups items into one mutable sequence. Empty square brackets create an empty list. Python permits duplicate elements and mixed types, as illustrated by a record containing a name, an age, and a boolean. This flexibility does not mean every operation will work on every combination of elements. For numerical score calculations, keeping the elements numeric is clearer. To update an existing item, we assign through an index. Replacing scores at index one changes seven to seven point five while keeping the list length unchanged. This is mutation of the list, unlike rebinding the name scores to a different list. An indexed assignment must target an existing position: assigning to index three in a three-item list does not append a new item. We use list methods to grow the list.''',
'Show the same three cells before and after mutation, with only the middle value changing. Do not illustrate replacement as inserting a fourth cell.')

slide('Adding items to a list', 75,
['append(x) adds one item at the end.', 'extend(iterable) adds its items individually.', 'insert(i, x) inserts an item before position i.', 'These methods mutate the list and return None.'],
'scores = [8.5, 7.0]\nscores.append(9.0)\nprint(scores)\nscores.extend([6.5, 8.0])\nprint(scores)\nscores.insert(1, 7.5)\nprint(scores)\nresult = scores.append(10.0)\nprint(result)',
'[8.5, 7.0, 9.0]\n[8.5, 7.0, 9.0, 6.5, 8.0]\n[8.5, 7.5, 7.0, 9.0, 6.5, 8.0]\nNone',
'''A method is an operation called on an object using a dot followed by the method name and parentheses. Append adds its argument as one item at the end. Extend takes items from an iterable and adds them individually; a list is one example of an iterable. Appending a list would add that entire list as one nested item, whereas extending with it adds its elements separately. Insert places an item before the specified position and shifts later items to the right. These methods modify the existing list and return None. Therefore, assigning scores equals scores dot append would lose the name's reference to the list and bind it to None. In our last example, result is None, but the list itself has still received the additional score.''',
'Reveal one operation at a time with its resulting list. Show the return value None separately from the changed list.')

slide('Removing items from a list', 70,
['remove(value) deletes the first matching value.', 'pop(index) removes and returns the item at that index.', 'pop() uses the last position by default.', 'del deletes by index or slice; clear() removes all items.'],
'scores = [8.5, 7.0, 9.0, 7.0]\nscores.remove(7.0)\nprint(scores)\nremoved = scores.pop()\nprint(removed)\nprint(scores)\ndel scores[0]\nprint(scores)\nscores.clear()\nprint(scores)',
'[8.5, 9.0, 7.0]\n7.0\n[8.5, 9.0]\n[9.0]\n[]',
'''Removal methods differ in how they identify an item and what they return. Remove searches by value and deletes the first equal item, so one of our two sevens remains. If the value is absent, remove raises a ValueError. Pop selects by position and returns the removed item, allowing us to store it in another variable. With no argument, it removes the last item. An invalid pop index, including popping from an empty list, raises an IndexError. Del is a statement rather than a method; it can delete an indexed item or a slice. Clear empties the existing list. Remove and clear return None, while pop returns the removed value. Choosing the correct operation depends on whether we know a value, know a position, or need the removed item afterward.''',
'Highlight value versus index in the method labels. Show the list after each operation so the changing indices remain clear.')

slide('Mutability and copying', 95,
['Lists can change in place; strings and tuples cannot have their items replaced.', 'b = a makes two names refer to the same list.', 'a.copy() creates a shallow copy of the outer list.', 'A tuple can contain a mutable object; immutability does not make all nested contents immutable.'],
'a = [8.5, 7.0]\nb = a\nb.append(9.0)\nprint(a)\nc = a.copy()\nc[0] = 10.0\nprint(a)\nprint(c)',
'[8.5, 7.0, 9.0]\n[8.5, 7.0, 9.0]\n[10.0, 7.0, 9.0]',
'''Mutability connects lists back to our model of variables. Assigning b equals a does not copy the list. Both names refer to the same object, so appending through b is visible when we print a. Copy creates a new outer list. Replacing its first element therefore leaves a unchanged. This is a shallow copy: if the elements were themselves mutable lists, those inner objects would still be shared. Our score list contains numbers, so replacing a score in the copy is independent. Strings and tuples do not support replacing their indexed elements. A variable referring to a string can still be reassigned; immutable describes the object, not a permanent restriction on the variable name. A tuple may also contain a mutable list whose contents can change, even though the tuple's own item references cannot be replaced.''',
'Use a reference diagram: a and b point to one list; c points to another. Include the shallow-copy caveat in a short caption.')

slide('Common mistakes and corrections', 75,
['Match operations to types; convert valid numeric text before arithmetic.', 'Use indices from 0 to len(sequence) - 1 for nonnegative indexing.', 'Do not assign the result of append back to the list variable.', 'Choose remove by value or pop by index deliberately.'],
'score = float("8.5")\nprint(score + 1)\nscores = [8.5, 7.0, 9.0]\nscores.append(8.0)\nprint(scores[-1])\nremoved = scores.pop(1)\nprint(removed)', '9.5\n8.0\n7.0',
'''Before combining the concepts, we can summarize several common errors. Adding the string eight point five to an integer raises a TypeError; converting valid numeric text fixes the mismatch. Int applied directly to the decimal string eight point five instead raises a ValueError, because the string is not an accepted integer representation. Accessing an item at len of the list is one position too far for a nonnegative index. The last item is at length minus one, or simply minus one. Assigning the return value of append to the list variable is another error: append changes the list but returns None. Finally, remove with the argument one searches for the value one, while pop with the argument one removes the second item. The corrected code demonstrates these distinctions without requiring audience responses.''',
'Use a compact mistake and correction table. Keep erroneous snippets separate from the runnable corrected code; no audience quiz or response pause.')

slide('Application data and updates', 95,
['Goal: store a student name and update a list of scores.', 'Convert a score supplied as text.', 'Preserve the original list with a shallow copy.', 'Assume valid numeric scores on a 0 to 10 scale and a nonempty list.'],
'student_name = "Alex"\nscores = [8.5, 7.0, 9.0]\noriginal_scores = scores.copy()\nnew_score_text = "8.0"\nnew_score = float(new_score_text)\nscores.append(new_score)\nscores[1] = 7.5\nprint(original_scores)\nprint(scores)',
'[8.5, 7.0, 9.0]\n[8.5, 7.5, 9.0, 8.0]',
'''We now apply the theory to a small student score program. Our goal is to store a student's name, add one score received as text, correct an existing score, and then calculate a summary. This example uses valid scores on a zero-to-ten scale and begins with a nonempty list. These are assumptions, not validation performed by the code. Student name is a string, and scores is a list of floats. We copy the list before updating it, allowing us to compare the original and revised values. We convert the new score text to a float, then append that number to the existing list. Finally, index one identifies the second score, which we correct to seven point five. The two printed lists confirm that the original copy remains unchanged. We can now calculate results using the updated data.''',
'Display the code in three groups: create and copy, convert and append, update and compare. Keep the original and revised lists visible side by side.')

slide('Application calculations and output', 100,
['sum(scores) calculates the total; len(scores) counts the scores.', 'Average = total / count; the list must be nonempty.', 'A comparison produces the boolean pass result.', 'Indexing and slicing retrieve selected scores.'],
'student_name = "Alex"\nscores = [8.5, 7.5, 9.0, 8.0]\naverage = sum(scores) / len(scores)\npassed = average >= 5.0\nprint("Student:", student_name)\nprint("Scores:", scores)\nprint("First score:", scores[0])\nprint("Last two:", scores[-2:])\nprint("Average:", average)\nprint("Passed:", passed)',
'Student: Alex\nScores: [8.5, 7.5, 9.0, 8.0]\nFirst score: 8.5\nLast two: [9.0, 8.0]\nAverage: 8.25\nPassed: True',
'''Using the updated list, sum produces a total of thirty-three and len produces a count of four. Dividing gives an average of eight point two five. This calculation requires at least one score; an empty list would make the denominator zero and raise a ZeroDivisionError. For this demonstration only, we define passing as an average of at least five. This is an example rule, not a claim about the course's grading policy. The comparison creates the boolean True. The output also retrieves the first score through indexing and the final two scores through slicing. In print, commas separate arguments, and Python displays them with spaces by default; we do not need to concatenate a string and a float ourselves. Together, the program demonstrates names, types, conversion, list mutation, copying, sequence access and a calculation. The following reference appendix contains the full program in one place.''',
'Show the six-line output next to the relevant calculations. The repeated data setup makes this slide runnable on its own; in the full demo, continue from the preceding slide.')

slide('Summary', 40,
['Variables are names bound to objects; objects have types.', 'Types determine valid operations and conversions.', 'Sequences support positional access; slicing excludes the stop bound.', 'Lists are mutable; assignment alone does not copy them.', 'Small programs combine these concepts to store and process data.'],
None, None,
'''To conclude, variables give us names for objects, and data types determine how those objects behave in expressions. Sequences organize data by position, making indexing and slicing possible. Lists add the ability to change their contents, but shared references require care because ordinary assignment does not create a copy. Our score example combines these ideas in a complete calculation with clearly stated assumptions. Understanding the relationship between names, types and mutable objects provides a foundation for later topics such as loops, functions and more complex data structures.''',
'Use five concise takeaway lines. End with a brief reference slide; no interactive quiz is needed.')

# Verify every runnable slide independently against its stated output.
for i,s in enumerate(slides,1):
    if s['code']:
        stream=io.StringIO()
        with contextlib.redirect_stdout(stream): exec(s['code'], {})
        assert stream.getvalue().strip()==s['output'].strip(), (i,stream.getvalue(),s['output'])
errors={'"8.5" + 1':TypeError, 'int("8.5")':ValueError, '[1, 2][2]':IndexError, '[1].pop(2)':IndexError, '[1].remove(2)':ValueError, 'sum([]) / len([])':ZeroDivisionError, '[1, 2][::0]':ValueError}
for code,kind in errors.items():
    try: eval(code)
    except kind: pass
    else: raise AssertionError(code)

doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.75)
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.05
doc.styles['Title'].font.size=Pt(27)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
code_style=doc.styles.add_style('Python Code',1)
code_style.font.name='Consolas'; code_style.font.size=Pt(9)
code_style.paragraph_format.space_after=Pt(0)
code_style.paragraph_format.line_spacing=1

def para(text,style=None): return doc.add_paragraph(text,style)
def codeblock(text):
    for line in text.split('\n'): para(line,'Python Code')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):
        c.text=h
        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E6EDF3'); c._tc.get_or_add_tcPr().append(shade)
        for r in c.paragraphs[0].runs: r.bold=True
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    for row in t.rows:
        for c in row.cells:
            pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(5); p.paragraph_format.space_before=Pt(5)
    return t

doc.add_heading('Python Variables and Data Types\nand Lists and Sequences',0)
para('Slide content and complete explanation script','Subtitle')
para('Introduction to Programming | English video assignment | Version 1')
para('This document provides a complete teaching sequence for the two assigned topics. It develops the theory first, supports each concept with a small example, and finishes with a student score application. The core lesson is that names refer to typed objects, while sequences organize items by position and lists allow those items to be changed.')
doc.add_heading('How to use this document',1)
para('Each numbered section represents one presentation slide. Slide content is the concise material to display. Code and output supply the worked example. Explanation is the full spoken content or speaker notes. Visual guidance describes how to arrange the material; it is not spoken narration. Do not place the full explanation paragraph on the slide.')
para('The plan contains 18 teaching slides plus a reference slide. Target recording length is approximately 20 to 22 minutes, including code walkthroughs and transitions. Timings are pacing guides rather than a requirement to fill silence. There are no audience-response pauses or live quizzes. The appendices support preparation and are not additional narrated slides.')
doc.add_heading('Lesson structure',1)
table(['Slides','Focus','Approximate time'],[['1','Introduction','0:40'],['2 to 6','Variables and basic types','5:45'],['7 to 10','Sequences and shared operations','4:40'],['11 to 15','Lists and common mistakes','6:20'],['16 to 17','Practical application','3:15'],['18','Summary','0:40']],[.8,3.9,1.9])
para('Preparation convention: all examples use Python 3. Each slide code block runs independently. For the continuous application recording, use the full program in Appendix A. Expected outputs have been checked by executing the examples. Intentional error examples are discussed separately so they do not interrupt the runnable blocks.')

for i,s in enumerate(slides,1):
    doc.add_page_break()
    doc.add_heading(f'{i:02d} {s["title"]}',1)
    para(f'Estimated pacing  {s["seconds"] // 60}:{s["seconds"] % 60:02d}')
    doc.add_heading('Slide content',2)
    for b in s['bullets']: para(b,'List Bullet')
    if s['code']:
        doc.add_heading('Code example',2); codeblock(s['code'])
        doc.add_heading('Expected output',2); codeblock(s['output'])
    doc.add_heading('Explanation',2); para(s['explanation'])
    doc.add_heading('Visual guidance',2); para(s['visual'])

doc.add_page_break(); doc.add_heading('Appendix A Complete application',1)
para('Run this program from top to bottom for the continuous demonstration. It uses a flat list of numbers, valid score text and at least one score. It does not implement user-input validation or a complete grading system.')
full='''student_name = "Alex"
scores = [8.5, 7.0, 9.0]
original_scores = scores.copy()

new_score_text = "8.0"
new_score = float(new_score_text)
scores.append(new_score)
scores[1] = 7.5

average = sum(scores) / len(scores)
passed = average >= 5.0

print("Student:", student_name)
print("Original scores:", original_scores)
print("Updated scores:", scores)
print("First score:", scores[0])
print("Last two:", scores[-2:])
print("Average:", average)
print("Passed:", passed)'''
codeblock(full)
stream=io.StringIO()
with contextlib.redirect_stdout(stream): exec(full,{})
doc.add_heading('Expected output',2); codeblock(stream.getvalue().strip())
doc.add_heading('Explanation of the assumptions',2)
para('The pass threshold of 5.0 is illustrative. Scores are assumed to be valid numeric values on a 0 to 10 scale. float() converts suitable numeric text but does not enforce this range. A production version would validate conversion, score bounds and the presence of at least one score before calculating the average. Those features require control flow and error handling beyond this lesson.')

doc.add_page_break(); doc.add_heading('Appendix B Review and presentation checklist',1)
doc.add_heading('Technical review',2)
for t in ['Keep assignment, equality, mutation and copying distinct.', 'Describe list order as positional order, not automatic sorting.', 'State that an immutable tuple can contain mutable objects.', 'Explain that list slicing and copy() are shallow operations.', 'Do not suggest that all sequence types support concatenation and repetition.', 'Keep remove(value), pop(index) and method return values accurate.', 'Run the final demo from a fresh session before recording.']:
    para(t,'List Bullet')
doc.add_heading('Slide and recording review',2)
for t in ['Use short on-screen points and place the Explanation text in speaker notes.', 'Reveal long code blocks in logical groups; make code readable at the final recording resolution.', 'Retain exact quotation marks, brackets and indentation when transferring code.', 'Use one consistent name and score example across the deck.', 'Record a short audio sample before recording the full presentation.', 'Rehearse the complete explanation once and adjust pacing naturally.', 'End with the summary and references rather than an audience quiz.']:
    para(t,'List Bullet')
doc.add_heading('Scope boundaries',2)
para('The lesson does not attempt to cover every Python type or every list method. Dictionaries, sets, complex arithmetic, comprehensions, loops, functions, deep copying and exception handling are outside the main sequence. This keeps the focus on the two assigned topics while retaining the caveats needed for accurate explanations.')

doc.add_page_break(); doc.add_heading('References',1)
para('Use a compact references slide at the end of the video. These official Python documentation pages support the definitions and behavior explained in the lesson. Examples and teaching explanations in this document are written for this assignment.')
refs=[('Python Tutorial An Informal Introduction to Python','https://docs.python.org/3/tutorial/introduction.html','Numbers, strings, assignment and the introduction to lists.'),('Python Tutorial Data Structures','https://docs.python.org/3/tutorial/datastructures.html','List methods, tuples and sequence concepts.'),('Python Standard Library Built in Types','https://docs.python.org/3/library/stdtypes.html','Sequence operations, mutability, ranges and method behavior.'),('Python Language Reference Data Model','https://docs.python.org/3/reference/datamodel.html','Objects, types, values and mutability.'),('Python Standard Library Built in Functions','https://docs.python.org/3/library/functions.html','type, input, int, float, str, len and sum.'),('Python Language Reference Lexical Analysis','https://docs.python.org/3/reference/lexical_analysis.html','Identifiers and keywords.')]
for title,url,use in refs:
    doc.add_heading(title,2); para(url); para(use)
para('Accessed 21 October 2026. The examples use long-established Python 3 behavior and were executed with Python 3.12 during preparation.')
footer=sec.footer.paragraphs[0]; footer.alignment=2
r=footer.add_run(); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); r._r.addnext(f)
doc.core_properties.title='Python Variables and Data Types and Lists and Sequences'
doc.core_properties.subject='Slide content and complete English explanation script'
doc.core_properties.author=''
out=ROOT/'Python_Lesson_Script.docx'; doc.save(out)
(ROOT/'slides.json').write_text(json.dumps(slides,indent=2),encoding='utf-8')
print(json.dumps({'file':str(out),'slides':len(slides),'pacing_seconds':sum(s['seconds'] for s in slides),'spoken_words':sum(len(s['explanation'].split()) for s in slides),'verified_code_blocks':sum(bool(s['code']) for s in slides),'verified_error_cases':len(errors)},indent=2))
